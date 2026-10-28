from django.contrib import messages
from .google_calendar_sync import sync_inspection
from datetime import datetime, date, time
from decimal import Decimal, InvalidOperation
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Q, Sum
from django.http import HttpResponseForbidden
from django.shortcuts import render, get_object_or_404, redirect
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from .models import Inspection, Membership


def company_queryset(user):
    return Inspection.objects.filter(company__memberships__user=user).distinct()


@login_required
def dashboard(request):
    qs = company_queryset(request.user)
    q = request.GET.get('q', '').strip()
    status = request.GET.get('status', '').strip()
    if q:
        qs = qs.filter(Q(job_number__icontains=q) | Q(insured_name__icontains=q) | Q(property_address__icontains=q) | Q(town__icontains=q))
    if status:
        qs = qs.filter(status=status)
    all_jobs = company_queryset(request.user)
    can_import = Membership.objects.filter(user=request.user, role__in=['owner', 'admin']).exists()
    return render(request, 'inspections/dashboard.html', {
        'jobs': qs[:200], 'q': q, 'status': status,
        'statuses': all_jobs.values_list('status', flat=True).distinct().order_by('status'),
        'total': all_jobs.count(),
        'outstanding': all_jobs.filter(paid=False, report_submitted_date__isnull=False).aggregate(value=Sum('job_fee'))['value'] or 0,
        'can_import': can_import,
    })


@login_required
def job_detail(request, pk):
    job = get_object_or_404(company_queryset(request.user), pk=pk)
    return render(request, 'inspections/detail.html', {'job': job})


def _txt(v):
    return str(v).strip() if v is not None else ''


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        try:
            return from_excel(v).date()
        except (ValueError, OverflowError, AttributeError):
            return None
    if isinstance(v, str) and v.strip():
        for fmt in ('%m/%d/%Y', '%Y-%m-%d', '%m/%d/%y'):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _as_datetime(v):
    if isinstance(v, datetime):
        d = v
    elif isinstance(v, (int, float)):
        try:
            d = from_excel(v)
        except (ValueError, OverflowError):
            return None
    elif isinstance(v, str) and v.strip():
        d = None
        for fmt in ('%m/%d/%Y %I:%M %p', '%Y-%m-%d %H:%M:%S', '%m/%d/%Y %H:%M'):
            try:
                d = datetime.strptime(v.strip(), fmt)
                break
            except ValueError:
                pass
        if d is None:
            return None
    else:
        return None
    return timezone.make_aware(d, timezone.get_current_timezone()) if timezone.is_naive(d) else d


def _fee(v):
    if v is None or _txt(v) == '':
        return None
    try:
        return Decimal(_txt(v).replace('$', '').replace(',', '')).quantize(Decimal('0.01'))
    except (InvalidOperation, ValueError):
        return None


@login_required
def import_cases(request):
    memberships = list(Membership.objects.filter(user=request.user, role__in=['owner', 'admin']).select_related('company'))
    if not memberships:
        return HttpResponseForbidden('Only a company owner or administrator can import cases.')
    if request.method == 'POST':
        company_id = request.POST.get('company_id', '')
        member = next((m for m in memberships if str(m.company_id) == company_id), None)
        if not member:
            messages.error(request, 'Select a company you administer.')
            return redirect('import_cases')
        upload = request.FILES.get('workbook')
        if not upload or not upload.name.lower().endswith('.xlsx'):
            messages.error(request, 'Choose an .xlsx Excel workbook.')
            return redirect('import_cases')
        if upload.size > 10 * 1024 * 1024:
            messages.error(request, 'The workbook must be 10 MB or smaller.')
            return redirect('import_cases')
        try:
            wb = load_workbook(upload, read_only=True, data_only=True)
            if 'Assignments' not in wb.sheetnames:
                wb.close()
                messages.error(request, 'This workbook has no Assignments tab.')
                return redirect('import_cases')
            ws = wb['Assignments']
            headers = [str(v or '').strip() for v in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
            if 'Job #' not in headers:
                wb.close()
                messages.error(request, 'Assignments row 2 must contain a Job # heading.')
                return redirect('import_cases')
            index = {name: i for i, name in enumerate(headers) if name}
            seen = set()
            created = existing = repeated = skipped = 0
            errors = []
            field_names = {
                'Insured Name': 'insured_name', 'Property Address': 'property_address',
                'Area / Calendar Group': 'area', 'Town': 'town', 'Phone': 'phone',
                'Policy #': 'policy_number', 'Insurance Carrier': 'insurance_carrier',
                'Inspection Company': 'inspection_company', 'Agent Name': 'agent_name',
                'Agent Phone': 'agent_phone', 'Contact Notes': 'contact_notes',
                'Notes': 'general_notes', 'Exterior Notes': 'exterior_notes',
                'Roof Notes': 'roof_notes', 'Electrical Notes': 'electrical_notes',
                'Plumbing Notes': 'plumbing_notes', 'HVAC Notes': 'hvac_notes',
                'Occupancy Notes': 'occupancy_notes', 'Hazards Notes': 'hazards_notes',
                'Photos Measurements': 'photos_measurements', 'General Field Notes': 'field_notes',
                'Calendar Event ID': 'calendar_event_id',
            }
            with transaction.atomic():
                for row_num, row in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
                    def val(name):
                        i = index.get(name)
                        return row[i] if i is not None and i < len(row) else None
                    job = _txt(val('Job #'))
                    if not job or job.upper().startswith('SAMPLE-'):
                        skipped += 1
                        continue
                    if job in seen:
                        repeated += 1
                        continue
                    seen.add(job)
                    if Inspection.objects.filter(company=member.company, job_number=job).exists():
                        existing += 1
                        continue
                    values = {model_field: _txt(val(column)) for column, model_field in field_names.items()}
                    values.update({
                        'date_assigned': _as_date(val('Date Assigned')),
                        'job_fee': _fee(val('Job Fee')),
                        'status': _txt(val('Status')) or 'New',
                        'scheduled_at': _as_datetime(val('Scheduled Date/Time')),
                        'inspection_complete_date': _as_date(val('Inspection Complete Date')),
                        'report_submitted_date': _as_date(val('Report Submitted Date')),
                        'payment_date': _as_date(val('Payment Date')),
                        'next_action_date': _as_date(val('Next Action Date')),
                        'paid': _txt(val('Paid?')).lower() in ('yes', 'true', '1'),
                    })
                    try:
                        with transaction.atomic():
                            Inspection.objects.create(company=member.company, job_number=job, **values)
                        created += 1
                    except Exception as exc:
                        errors.append(f'Row {row_num} ({job}): {exc}')
            wb.close()
            return render(request, 'inspections/import_cases.html', {
                'memberships': memberships,
                'result': {'created': created, 'existing': existing, 'repeated': repeated, 'skipped': skipped, 'errors': errors[:20], 'error_count': len(errors)},
            })
        except Exception as exc:
            messages.error(request, f'Unable to read workbook: {exc}')
            return redirect('import_cases')
    return render(request, 'inspections/import_cases.html', {'memberships': memberships})
from django import forms
from django.contrib import messages
from django.shortcuts import redirect
from django.contrib.auth.decorators import login_required

EDIT_FIELDS = [
    'insured_name',
    'property_address',
    'phone',
    'status',
    'scheduled_at',
    'job_fee',
    'contact_notes',
    'general_notes',
    'inspection_complete_date',
    'report_submitted_date',
    'paid',
    'payment_date',
]

class InspectionEditForm(forms.ModelForm):
    class Meta:
        model = Inspection
        fields = EDIT_FIELDS
        widgets = {
            'scheduled_at': forms.DateTimeInput(
                attrs={'type': 'datetime-local'},
                format='%Y-%m-%dT%H:%M',
            ),
            'inspection_complete_date': forms.DateInput(
                attrs={'type': 'date'},
                format='%Y-%m-%d',
            ),
            'report_submitted_date': forms.DateInput(
                attrs={'type': 'date'},
                format='%Y-%m-%d',
            ),
            'payment_date': forms.DateInput(
                attrs={'type': 'date'},
                format='%Y-%m-%d',
            ),
        }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields['scheduled_at'].input_formats = ['%Y-%m-%dT%H:%M']
        for field in ('inspection_complete_date', 'report_submitted_date', 'payment_date'):
            self.fields[field].input_formats = ['%Y-%m-%d']



@login_required
def job_edit(request, pk):
    job = get_object_or_404(company_queryset(request.user), pk=pk)

    if request.method == 'POST':
        form = InspectionEditForm(request.POST, instance=job)

        if form.is_valid():
            job = form.save()

            if job.scheduled_at:
                try:
                    sync_inspection(job)
                    messages.success(
                        request,
                        'Inspection saved and Google Calendar updated.'
                    )
                except Exception as error:
                    messages.warning(
                        request,
                        f'Inspection saved, but Calendar sync failed: {error}'
                    )
            else:
                messages.success(request, 'Inspection saved.')

            return redirect('job_detail', pk=job.pk)

    else:
        form = InspectionEditForm(instance=job)

    return render(request, 'inspections/edit.html', {
        'job': job,
        'form': form,
    })


import calendar
from datetime import date, datetime, time, timedelta
from django.utils import timezone


@login_required
def calendar_view(request):
    today = timezone.localdate()

    try:
        year = int(request.GET.get("year", today.year))
        month = int(request.GET.get("month", today.month))
        first_day = date(year, month, 1)
    except (ValueError, TypeError):
        first_day = today.replace(day=1)

    year = first_day.year
    month = first_day.month

    calendar_data = calendar.Calendar(firstweekday=6)
    weeks = calendar_data.monthdatescalendar(year, month)

    start = weeks[0][0]
    end = weeks[-1][-1] + timedelta(days=1)

    jobs = company_queryset(request.user).filter(
        scheduled_at__gte=timezone.make_aware(
            datetime.combine(start, time.min)
        ),
        scheduled_at__lt=timezone.make_aware(
            datetime.combine(end, time.min)
        ),
    ).order_by("scheduled_at")

    jobs_by_day = {}

    for job in jobs:
        day = timezone.localtime(job.scheduled_at).date()
        jobs_by_day.setdefault(day, []).append(job)

    calendar_weeks = [
        [
            {
                "date": day,
                "current_month": day.month == month,
                "jobs": jobs_by_day.get(day, []),
            }
            for day in week
        ]
        for week in weeks
    ]

    previous_month = first_day - timedelta(days=1)
    next_month = (first_day.replace(day=28) + timedelta(days=4)).replace(day=1)

    return render(request, "inspections/calendar.html", {
        "calendar_weeks": calendar_weeks,
        "month_title": first_day.strftime("%B %Y"),
        "previous_month": previous_month,
        "next_month": next_month,
        "today": today,
    })


from django.utils import timezone


@login_required
def worklist_view(request):
    today = timezone.localdate()
    jobs = company_queryset(request.user)

    todays_appointments = jobs.filter(
        scheduled_at__date=today
    ).order_by("scheduled_at")

    calls_to_make = jobs.filter(
        status__in=[
            "New",
            "Call #1",
            "Call #2",
            "Agent Call",
        ],
        scheduled_at__isnull=True,
    ).order_by("date_assigned", "job_number")

    overdue_followups = jobs.filter(
        next_action_date__lt=today,
        scheduled_at__isnull=True,
    ).exclude(
        status__in=["Submitted", "Paid", "Cancelled"]
    ).order_by("next_action_date")

    reports_to_submit = jobs.filter(
        inspection_complete_date__isnull=False,
        report_submitted_date__isnull=True,
    ).exclude(
        status__in=["Submitted", "Paid", "Cancelled"]
    ).order_by("inspection_complete_date")

    payments_outstanding = jobs.filter(
        report_submitted_date__isnull=False,
        paid=False,
    ).exclude(
        status__in=["Paid", "Cancelled"]
    ).order_by("report_submitted_date")

    return render(request, "inspections/worklist.html", {
        "today": today,
        "todays_appointments": todays_appointments,
        "calls_to_make": calls_to_make,
        "overdue_followups": overdue_followups,
        "reports_to_submit": reports_to_submit,
        "payments_outstanding": payments_outstanding,
    })


from datetime import date
from django.views.decorators.http import require_POST
from django.shortcuts import redirect, get_object_or_404

@login_required
@require_POST
def quick_action(request, pk):
    job = get_object_or_404(
        company_queryset(request.user),
        pk=pk
    )

    action = request.POST.get("action")

    if action == "status":
        new_status = request.POST.get("status", "").strip()

        allowed_statuses = [
            "New", "Call #1", "Call #2", "Agent Call",
            "Scheduled", "Submitted", "Paid", "Cancelled",
        ]

        if new_status in allowed_statuses:
            job.status = new_status
            job.save(update_fields=["status"])

        return redirect("worklist")

    if action == "followup":
        followup_date = request.POST.get("followup_date")

        if followup_date:
            try:
                job.next_action_date = date.fromisoformat(followup_date)
            except ValueError:
                return redirect("worklist")
        else:
            job.next_action_date = None

        job.save(update_fields=["next_action_date"])

    return redirect("worklist")
