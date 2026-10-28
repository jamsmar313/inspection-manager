from django.core.management.base import BaseCommand,CommandError
from openpyxl import load_workbook
from collections import Counter
class Command(BaseCommand):
 help='Preview source workbook assignments without writing to database'
 def add_arguments(self,parser):parser.add_argument('workbook')
 def handle(self,*args,**opts):
  wb=load_workbook(opts['workbook'],read_only=True,data_only=True)
  if 'Assignments' not in wb:raise CommandError('Assignments tab not found')
  rows=wb['Assignments'].iter_rows(values_only=True)
  next(rows,None);headers=[str(x or '').strip() for x in next(rows)]
  if 'Job #' not in headers:raise CommandError('Job # column missing')
  idx=headers.index('Job #');ids=[str(row[idx]).strip() for row in rows if len(row)>idx and row[idx] is not None and str(row[idx]).strip()]
  duplicates=[k for k,v in Counter(ids).items() if v>1]
  self.stdout.write(f'Assignment records: {len(ids)}; unique job numbers: {len(set(ids))}; repeated job numbers: {len(duplicates)}')
  for k in duplicates[:30]:self.stdout.write('DUPLICATE: '+k)
  self.stdout.write('No records were changed. Review duplicates before importing.')
